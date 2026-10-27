#!/usr/bin/env python3
"""Build the idea DOCX/PDF and commercial/evaluation PDFs from verified results."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    HRFlowable, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph,
    SimpleDocTemplate, Spacer, Table, TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
NAVY = "142C43"
TEAL = "16827E"
SLATE = "536A79"
PALE = "EFF5F6"
GREEN = "E5F2EC"
AMBER = "FBF3E6"
RED = "F8EBE9"


def load_optional(path: Path, default):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def fmt_seconds(value: float | None) -> str:
    return "not measured" if value is None else f"{value * 1000:.3f} ms"


def fmt_pct(item: dict) -> str:
    return f"{item['passed']}/{item['total']} ({item['rate'] * 100:.1f}%)"


evaluation = load_optional(ART / "evaluation-results.json", {})
pytest_summary = evaluation.get("test_suite", {"passed": "not recorded", "failed": "not recorded", "total": "not recorded"})
checks = {item["name"]: item for item in evaluation.get("checks", [])}
restore_stats = evaluation.get("restore_duration_seconds", {})
security = load_optional(ART / "security-results.json", {})


def add_word_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def shade_cell(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); props.append(shd)


def set_cell_padding(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}"); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); tcMar.append(node)
    tcPr.append(tcMar)


idea_pages = [
    {
        "kicker": "01 / THE CASE",
        "title": "The recovery gap",
        "paragraphs": [
            "Critical services depend on data, identity, networks, power, people and facilities. A backup job can be green while its copy is stale, altered, trapped in the failed site, dependent on unavailable identity, or never successfully restored. During a ransomware event or physical disaster, teams need to establish what can return, in what order, and what evidence supports that conclusion.",
            "ResilienceGuard turns that question into a repeatable readiness assessment. It joins inventory, backup metadata, failure domains, dependency records, recovery objectives and isolated test evidence. The platform is designed to complement established backup products: it assesses recoverability across sources rather than storing enterprise data itself.",
        ],
        "bullets": [
            "Decision makers: CIO/CISO, continuity and disaster recovery owners, infrastructure leads, risk and assurance teams.",
            "Initial users: hospitals, utilities, public services and critical suppliers with small or fragmented recovery teams.",
            "Outcome: a defensible view of available recovery points, blockers, next steps and observed test evidence.",
        ],
        "callout": "Working MVP: four synthetic services, five backup records, three disruption scenarios, deterministic dependency analysis, human-gated isolated file restoration, measured timing, and PDF evidence export.",
    },
    {
        "kicker": "02 / NEED & EVIDENCE",
        "title": "Why recovery assurance matters",
        "paragraphs": [
            "Backup readiness is a lifecycle: make copies, protect and locate them, validate integrity, rehearse restoration, and record the result. RPO and RTO targets only become useful when they are compared with an eligible recovery point and measured recovery exercise. Data center loss and cyber incidents can overlap; the dependencies that normally provide access to backups may also be unavailable.",
            "NIST SP 1339, finalized in June 2026, describes backups as important to OT recovery from reliability and cyber incidents and recommends recurring restore tests on non-production systems. NIST SP 800-34 Rev. 1 provides contingency-planning guidance for prioritization and recovery strategies. CSF 2.0 includes a Recover function. NATO's public 2026 resilience baseline emphasizes essential services, physical and cyber risk, cascading impacts and recovery capability. These sources inform the problem framing; the product does not certify alignment.",
        ],
        "bullets": [
            "NIST SP 1339: https://csrc.nist.gov/pubs/sp/1339/final",
            "NIST SP 800-34 Rev. 1: https://csrc.nist.gov/pubs/sp/800/34/r1/upd1/final",
            "NIST CSF 2.0: https://www.nist.gov/cyberframework",
            "NATO 2026 Baseline Requirements: https://www.nato.int/en/about-us/official-texts-and-resources/official-texts/2026/09/15/2026-resilience-baseline-requirements",
        ],
        "callout": "The product's evidence model keeps backup existence, freshness, hash integrity, malware status and actual restore success as separate states.",
    },
    {
        "kicker": "03 / PRODUCT",
        "title": "From evidence to a safe next test",
        "paragraphs": [
            "An operator chooses a disaster scenario and sees a dashboard of service availability, backup status, failure-domain reachability, dependency blockers and RPO/RTO gaps. The deterministic engine proposes a recovery order. A reviewer can create a sandbox-only restore proposal for a registered synthetic file, inspect its evidence and explicitly approve or reject it. Approval copies the file into a dedicated isolated output directory and records the measured copy/check time, SHA-256 comparison and fixture health result.",
            "The assistant can summarize the evidence and surface uncertainties. Its optional local model receives a bounded facts package and a permitted action/evidence list. The model cannot call tools or execute restoration. If it is absent or returns invalid output, a deterministic summary is used. The approval workflow is persisted and auditable.",
        ],
        "bullets": [
            "Import: service catalog, backup metadata, manifest digest, location, recovery target and dependencies.",
            "Assess: freshness, integrity, location reachability, scenario trust and dependency readiness.",
            "Propose: dependency-aware sequence plus a list of uncertainties and tests to perform.",
            "Validate: human decision, isolated sample restore, checksum comparison, basic fixture parsing and report download.",
        ],
        "callout": "No production connector, credential, backup-control API, destructive model tool or production write path is included.",
    },
    {
        "kicker": "04 / INVENTORY & SCENARIOS",
        "title": "A small but meaningful recovery graph",
        "paragraphs": [
            "The repeatable data generator uses seed 20261026 and an explicit reference time of 2026-10-26 08:00 UTC. It creates a Hospital Care Coordination application, Clinical Records Database, Clinical Identity Service and Directory Data Store. Dependencies are explicit: the application relies on the clinical database and identity; identity depends on directory data; recovery tasks require an independent recovery network and storage. The directory service intentionally has no backup record to test cascade detection.",
            "The primary-data-center scenario removes the primary facility and East vault while leaving the West vault and recovery site available. The ransomware scenario marks the newest database candidate corrupted and untrusted. The combined power/network scenario removes the East vault and primary access path while preserving the independent recovery site. The engine identifies the verified West database candidate as stale against its two-hour RPO and blocks the application through unavailable identity/directory dependencies.",
        ],
        "bullets": [
            "Five backup records across two locations, one intentionally altered snapshot and one intentionally missing service backup.",
            "Service-specific RPO and RTO targets, estimated restore times, retention metadata and stable evidence IDs.",
            "NetworkX DAG validation rejects unknown edges and cycles; deterministic traversal exposes cascading blockers.",
        ],
        "callout": "Synthetic findings are designed for repeatable demonstration; they are not a benchmark of a real hospital or operator.",
    },
    {
        "kicker": "05 / AI & TRUST",
        "title": "Governed assistance, bounded authority",
        "paragraphs": [
            "Safety-critical readiness and operation eligibility are computed by deterministic Python rules. SHA-256 detects whether the sample file matches the expected digest. It cannot establish who created the file or whether it is malware-free. The system displays malware cleanliness as unknown because no scanner runs. Actual restore success is recorded only after a real copy and validation.",
            "An optional Ollama 0.40.11 runtime with the Qwen3 8B Apache-2.0 model can produce an advisory summary. The prompt permits only supplied findings and existing action IDs. Responses are schema-checked and unknown evidence/action IDs are rejected. Natural-language claims can still be wrong, so model output is labeled advisory and asks for human review. Default mode is deterministic and works offline.",
        ],
        "bullets": [
            "No LangGraph: one linear approval state machine is persisted with SQLAlchemy and does not need agent checkpointing.",
            "Pending actions cannot be executed. Approved operations are fixed to a fixture copy under RESTORE_ROOT.",
            "Demo scope has no login, role separation, quorum approval, antivirus or immutable storage; operational deployment needs these controls.",
        ],
        "callout": "A hash match is not a clean-data verdict. The MVP restores only synthetic JSON and does not start a database server or application.",
    },
    {
        "kicker": "06 / IMPLEMENTATION",
        "title": "Local-first, deployable foundation",
        "paragraphs": [
            "The API uses FastAPI, Pydantic, SQLAlchemy and PostgreSQL 18 in Docker Compose, with SQLite for direct single-user local operation. React and TypeScript power the dashboard and workflow. NetworkX evaluates dependencies. The fixture source is mounted read-only; a separate volume holds isolated restore output. ReportLab creates evidence PDFs. PptxGenJS and python-docx generate the hackathon submission materials.",
            "Technology versions are pinned and dependency audits are run. The initial deployment intentionally avoids a distributed object store because local files are enough for the synthetic dataset. MinIO upstream was archived in April 2026; a production object-storage connector should evaluate current maintained alternatives and validate API compatibility rather than adding an unneeded service to the MVP.",
        ],
        "bullets": [
            "Compose starts Postgres, API and UI on localhost; fixture volume is read-only and restore data is isolated.",
            "No actual customer, patient, military or production data is required.",
            "The optional LLM profile is separate because model downloads, memory and hardware requirements vary.",
        ],
        "callout": "Architecture is designed for adapter-based expansion; vendor integrations, cloud controls and operations hardening are future work.",
    },
    {
        "kicker": "07 / EVALUATION",
        "title": "Measured synthetic results",
        "paragraphs": [
            f"The submitted automated suite recorded {pytest_summary.get('passed', 'not recorded')} passing tests and {pytest_summary.get('failed', 'not recorded')} failures. The synthetic evaluation checked {fmt_pct(checks.get('SHA-256 expected good/bad classification', {'passed': 0, 'total': 0, 'rate': 0}))} expected SHA-256 classifications, {fmt_pct(checks.get('RPO freshness classification', {'passed': 0, 'total': 0, 'rate': 0}))} freshness classifications, and {fmt_pct(checks.get('Recovery dependency matrix (4 services × 3 scenarios)', {'passed': 0, 'total': 0, 'rate': 0}))} service/scenario dependency labels.",
            f"Across {restore_stats.get('samples', 0)} actual local sandbox file-copy runs, mean measured copy/hash/fixture-check time was {fmt_seconds(restore_stats.get('mean'))}. These small synthetic file results are not application RTO. The evaluation compared the rule-only baseline with the deterministic assistant fallback; allowed action IDs matched exactly. No Ollama/Qwen3 inference was run, so model accuracy is not measured.",
        ],
        "bullets": [
            f"Approved isolated copy + file hash + fixture validation: {fmt_pct(checks.get('Approved isolated restore copy + hash + fixture health', {'passed': 0, 'total': 0, 'rate': 0}))}.",
            f"Unauthorized/untrusted restore prevention: {fmt_pct(checks.get('Unauthorized/untrusted restore prevention', {'passed': 0, 'total': 0, 'rate': 0}))}.",
            f"Evidence/action grounding against rule-only baseline: {fmt_pct(checks.get('Evidence/action grounding against rule-only baseline', {'passed': 0, 'total': 0, 'rate': 0}))}.",
            "All findings are synthetic and repeatable; no real data center or threat environment was exercised.",
        ],
        "callout": "Measured only after execution. Local generative-model accuracy and actual malware cleanliness remain untested.",
    },
    {
        "kicker": "08 / CUSTOMER & COMMERCIAL",
        "title": "A readiness layer that complements existing tools",
        "paragraphs": [
            "Economic buyers may include a CIO, CISO, COO, chief risk officer or continuity leader. Operational users include backup administrators, infrastructure teams, security incident responders and auditors. Deployment options include customer-hosted on-premises/private cloud, a managed isolated service, or a partner-delivered assessment. A pilot can normalize one service catalog, model critical dependencies, configure a small scenario set and run a supervised recovery exercise.",
            "Rubrik, Veeam and Cohesity already sell recovery orchestration, isolated testing, clean-room workflows or related security capabilities. ResilienceGuard's proposed positioning is vendor-independent evidence aggregation across existing backup systems and simultaneous cyber/physical disruption. It is not a replacement or a claim of capability superiority. Vendor interoperability and buyer demand need discovery interviews and connector pilots.",
        ],
        "bullets": [
            "Illustrative first-year per-customer assumptions: $18,000 implementation + $48,000 annual license + $12,000 annual support.",
            "Illustrative direct delivery/support cost: $18,000 per customer; recurring fees and support drive follow-on revenue.",
            "Potential revenue: readiness assessment, integration/implementation, annual subscription, annual support and exercises.",
        ],
        "callout": "The financial case is a planning model only. Prices, delivery effort, overhead, customer willingness to pay and margin have not been market-verified.",
    },
    {
        "kicker": "09 / RISKS & NEXT",
        "title": "Build trust through evidence",
        "paragraphs": [
            "The most serious risk is a readiness score that overstates safety. Incomplete dependency data can hide a blocker; a digest does not prove authenticity or malware absence; optional model prose can make unsupported claims; a local demo approval is not production identity. The MVP counters those limits with deterministic gates, explicit unknown states, read-only fixtures, confined output paths, a human decision and visible synthetic labels.",
            "Next work should prioritize authenticated users and independent approvals, signed manifests backed by external keys, antivirus/EDR integration in an isolated lab, schema adapters for existing backup products, more complete recovery dependencies, actual app/database health probes and an OpenTelemetry export with a reviewed privacy model. Pilot validation should compare assessments against existing tabletop/recovery exercises and track observed RPO/RTO outcomes over time.",
        ],
        "bullets": [
            "Phase 1: interviews and a supervised hospital/utility pilot with synthetic or de-identified inventories.",
            "Phase 2: vendor-neutral ingestion, role-based approval, deployment profiles and customer-owned evidence retention.",
            "Phase 3: validated clean-room integrations, exercise scheduling, aggregate reporting and managed support.",
        ],
        "callout": "ResilienceGuard's promise is not that an outage cannot happen. It is that teams can see recovery evidence, blockers and uncertainty before the outage does.",
    },
]


def make_idea_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.67); section.bottom_margin = Inches(.62); section.left_margin = Inches(.73); section.right_margin = Inches(.73)
    section.header_distance = Inches(.32); section.footer_distance = Inches(.32)
    normal = doc.styles["Normal"]; normal.font.name = "Arial"; normal.font.size = Pt(9.3); normal.font.color.rgb = RGBColor(45, 65, 78)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 30, NAVY), ("Heading 1", 19, NAVY), ("Heading 2", 12, TEAL)):
        style = doc.styles[name]; style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(2); style.paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]; header.text = "RESILIENCEGUARD  /  EUREKA HACKATHON"; header.style = doc.styles["Caption"]
    header.runs[0].font.name = "Arial"; header.runs[0].font.size = Pt(7.5); header.runs[0].font.bold = True; header.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("SYNTHETIC MVP  ·  "); add_word_field(footer, "PAGE")
    for run in footer.runs: run.font.size = Pt(7); run.font.color.rgb = RGBColor.from_string("8A9AA3")

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(55); p.paragraph_format.space_after = Pt(8)
    run = p.add_run("EUREKA HACKATHON  /  IDEA DOCUMENT"); run.bold = True; run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string(TEAL)
    title = doc.add_paragraph(style="Title"); title.paragraph_format.space_after = Pt(6); title.add_run("ResilienceGuard")
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(20)
    r = p.add_run("Governed Agentic AI for Cyber and Disaster Recovery Readiness"); r.font.size = Pt(16); r.font.color.rgb = RGBColor.from_string(SLATE)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(15)
    r = p.add_run("Know what comes back. Prove it in an isolated test."); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph("A lightweight, vendor-independent recovery-readiness and evidence layer for operators of essential services. It checks backup evidence and system dependencies, then lets a person approve a constrained copy into an isolated test folder. The optional AI explains findings; fixed rules decide eligibility.")
    cover_table = doc.add_table(rows=2, cols=2); cover_table.autofit = False
    cover_cells = [("THE QUESTION", "Can our most important services actually be restored after cyber or physical loss?"), ("THE ANSWER", "Deterministic evidence checks, dependency-aware plans and real isolated sample restores."), ("THE USERS", "Hospitals · utilities · government · critical suppliers"), ("THE MVP", "FastAPI · React · PostgreSQL · NetworkX · local advisory model option")]
    for idx, (head, body) in enumerate(cover_cells):
        cell = cover_table.cell(idx // 2, idx % 2); cell.width = Inches(3.15); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade_cell(cell, PALE); set_cell_padding(cell, 150, 170, 150, 170)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(4); r = p.add_run(head); r.bold = True; r.font.size = Pt(7); r.font.color.rgb = RGBColor.from_string(TEAL)
        p = cell.add_paragraph(body); p.paragraph_format.space_after = Pt(0); p.runs[0].font.size = Pt(9)
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("A quick example  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run("At 08:00, the 05:00 database copy is one hour older than its two-hour data-loss target. It matches its saved fingerprint, but the hospital app is blocked because a required directory backup is missing. No malware scan or live app restore has occurred.")
    p = doc.add_paragraph("For definitions of fingerprint, data-loss target, dependency, and other terms, see AUDIENCE_GUIDE.md.")
    p.paragraph_format.space_after = Pt(0); p.runs[0].italic = True; p.runs[0].font.size = Pt(8.5); p.runs[0].font.color.rgb = RGBColor.from_string(SLATE)
    status = doc.add_table(rows=1, cols=1); cell = status.cell(0, 0); shade_cell(cell, GREEN); set_cell_padding(cell, 140, 170, 140, 170)
    p = cell.paragraphs[0]; r = p.add_run("IMPLEMENTATION STATUS  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(TEAL); r.font.size = Pt(8)
    p.add_run(f"Synthetic MVP implemented. Automated checks: {pytest_summary.get('passed', 'not recorded')} passed; actual sandbox copy checks: {fmt_pct(checks.get('Approved isolated restore copy + hash + fixture health', {'passed': 0, 'total': 0, 'rate': 0}))}.")
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(27)
    r = p.add_run("Submission snapshot: 26 October 2026  ·  Version 0.1.0"); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string("83939C")

    for page in idea_pages:
        doc.add_page_break()
        kicker = doc.add_paragraph(); kicker.paragraph_format.space_before = Pt(2); kicker.paragraph_format.space_after = Pt(5)
        run = kicker.add_run(page["kicker"]); run.bold = True; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(TEAL)
        doc.add_heading(page["title"], level=1)
        for text in page["paragraphs"]:
            para = doc.add_paragraph(text); para.paragraph_format.keep_together = True
        for item in page["bullets"]:
            para = doc.add_paragraph(style="List Bullet"); para.paragraph_format.left_indent = Inches(.18); para.paragraph_format.first_line_indent = Inches(-.14); para.paragraph_format.space_after = Pt(4); para.add_run(item)
        callout = doc.add_table(rows=1, cols=1); cell = callout.cell(0, 0); shade_cell(cell, AMBER if page["kicker"].startswith("05") else PALE); set_cell_padding(cell, 135, 160, 135, 160)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run(page["callout"]); r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = RGBColor.from_string(NAVY)

    doc.add_heading("References and assumptions", level=2)
    refs = [
        "NIST SP 1339, OT Backup Quick Start Guide, June 2026. https://csrc.nist.gov/pubs/sp/1339/final",
        "NIST SP 800-34 Rev. 1, Contingency Planning Guide for Federal Information Systems. https://csrc.nist.gov/pubs/sp/800/34/r1/upd1/final",
        "NIST Cybersecurity Framework 2.0. https://www.nist.gov/cyberframework",
        "NATO 2026 Resilience Baseline Requirements, 15 September 2026. https://www.nato.int/en/about-us/official-texts-and-resources/official-texts/2026/09/15/2026-resilience-baseline-requirements",
        "Rubrik Cyber Recovery. https://www.rubrik.com/products/cyber-recovery",
        "Veeam Recovery Orchestrator overview. https://helpcenter.veeam.com/docs/vro/userguide/overview.html",
        "Cohesity Clean Room. https://www.cohesity.com/solutions/clean-room/",
        "MinIO upstream repository status. https://github.com/minio/minio",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4); run = p.add_run(ref); run.font.size = Pt(8.5)
    doc.add_heading("Interpretation", level=2)
    doc.add_paragraph("Standards and vendor pages establish context, not proof of product compliance or competitive superiority. The adoption, customer value, price, implementation effort and financial scenarios are hypotheses for validation. Any actual service deployment requires threat modeling, privacy/security review, customer-owned recovery policies and integrations tested in a representative isolated environment.")
    doc.add_heading("Illustrative first-year unit assumptions", level=2)
    doc.add_paragraph("Implementation $18,000 + annual license $48,000 + annual support $12,000; direct delivery/support cost $18,000 per customer. All figures are planning assumptions, not market-verified prices or quotes.")
    doc.save(ROOT / "ResilienceGuard_Idea_Document.docx")


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="RGTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=30, leading=34, textColor=colors.HexColor("#" + NAVY), alignment=TA_LEFT, spaceAfter=10))
    styles.add(ParagraphStyle(name="RGH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=19, leading=22, textColor=colors.HexColor("#" + NAVY), spaceBefore=2, spaceAfter=9))
    styles.add(ParagraphStyle(name="RGH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.HexColor("#" + TEAL), spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle(name="RGBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.1, leading=12.2, textColor=colors.HexColor("#2D414E"), spaceAfter=7))
    styles.add(ParagraphStyle(name="RGSmall", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.7, leading=10.1, textColor=colors.HexColor("#536A79"), spaceAfter=4))
    styles.add(ParagraphStyle(name="RGKicker", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.HexColor("#" + TEAL), spaceAfter=5))
    styles.add(ParagraphStyle(name="RGCallout", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=8.6, leading=11.5, textColor=colors.HexColor("#" + NAVY)))
    styles.add(ParagraphStyle(name="RGCentered", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=13, alignment=TA_CENTER, textColor=colors.HexColor("#536A79")))
    return styles


STYLES = pdf_styles()


def footer(canvas, doc):
    canvas.saveState(); w, h = letter
    canvas.setStrokeColor(colors.HexColor("#DDE6E9")); canvas.setLineWidth(.5); canvas.line(.65 * inch, .48 * inch, w - .65 * inch, .48 * inch)
    canvas.setFont("Helvetica", 7); canvas.setFillColor(colors.HexColor("#84949D")); canvas.drawString(.68 * inch, .32 * inch, "RESILIENCEGUARD  /  SYNTHETIC MVP")
    canvas.drawRightString(w - .68 * inch, .32 * inch, f"{doc.page}")
    canvas.setFillColor(colors.HexColor("#16827E")); canvas.setFont("Helvetica-Bold", 7); canvas.drawString(.68 * inch, h - .36 * inch, "EUREKA HACKATHON  /  RECOVERY READINESS")
    canvas.restoreState()


def pdf_bullet(items):
    return ListFlowable([ListItem(Paragraph(item, STYLES["RGBody"]), leftIndent=8) for item in items], bulletType="bullet", start="circle", leftIndent=14, bulletFontName="Helvetica", bulletFontSize=6, bulletColor=colors.HexColor("#" + TEAL), spaceAfter=5)


def pdf_callout(text, fill=PALE):
    table = Table([[Paragraph(text, STYLES["RGCallout"])]], colWidths=[7.15 * inch])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + fill)), ("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#D9E6E8")), ("LEFTPADDING", (0, 0), (-1, -1), 11), ("RIGHTPADDING", (0, 0), (-1, -1), 11), ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
    return table


def make_idea_pdf():
    out = ROOT / "ResilienceGuard_Idea_Document.pdf"
    doc = SimpleDocTemplate(str(out), pagesize=letter, rightMargin=.67 * inch, leftMargin=.67 * inch, topMargin=.62 * inch, bottomMargin=.62 * inch, title="ResilienceGuard Idea Document", author="ResilienceGuard")
    story = [Spacer(1, .9 * inch), Paragraph("EUREKA HACKATHON  /  IDEA DOCUMENT", STYLES["RGKicker"]), Paragraph("ResilienceGuard", STYLES["RGTitle"]), Paragraph("Governed Agentic AI for Cyber and Disaster Recovery Readiness", STYLES["RGH2"]), Spacer(1, 7), Paragraph("Know what comes back. Prove it in an isolated test.", ParagraphStyle(name="CoverLine", parent=STYLES["RGH1"], fontSize=17, leading=22)), Paragraph("A lightweight, vendor-independent recovery-readiness and evidence layer for operators of essential services. It checks backup evidence and system dependencies, then lets a person approve a constrained copy into an isolated test folder. The optional AI explains findings; fixed rules decide eligibility.", STYLES["RGBody"]), Spacer(1, 7)]
    cover_rows = [["THE QUESTION", "THE ANSWER"], ["Can our most important services be restored after cyber or physical loss?", "Deterministic checks, dependency-aware plans and real isolated sample restores."], ["THE USERS", "THE MVP"], ["Hospitals · utilities · government · critical suppliers", "FastAPI · React · PostgreSQL · NetworkX · local advisory model option"]]
    t = Table([[Paragraph(f"<b>{cell}</b>", STYLES["RGSmall"]) for cell in row] for row in cover_rows], colWidths=[3.575 * inch, 3.575 * inch])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EFF5F6")), ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#" + NAVY)), ("GRID", (0, 0), (-1, -1), .5, colors.white), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)])); story += [t, Spacer(1, 9), Paragraph("<b>A quick example:</b> At 08:00, the 05:00 database copy is one hour older than its two-hour data-loss target. Its fingerprint matches, but the hospital app is blocked because a required directory backup is missing. No malware scan or live app restore has occurred.", STYLES["RGBody"]), Paragraph("For plain-language definitions and the full example, see AUDIENCE_GUIDE.md.", STYLES["RGSmall"]), Spacer(1, 7), pdf_callout(f"IMPLEMENTATION STATUS  ·  Synthetic MVP implemented. Automated checks: {pytest_summary.get('passed', 'not recorded')} passed; actual sandbox copy checks: {fmt_pct(checks.get('Approved isolated restore copy + hash + fixture health', {'passed': 0, 'total': 0, 'rate': 0}))}.", GREEN), Spacer(1, 14), Paragraph("Submission snapshot: 26 October 2026  ·  Version 0.1.0", STYLES["RGSmall"])]
    for page in idea_pages:
        story.extend([PageBreak(), Paragraph(page["kicker"], STYLES["RGKicker"]), Paragraph(page["title"], STYLES["RGH1"])])
        story.extend(Paragraph(text, STYLES["RGBody"]) for text in page["paragraphs"])
        story.append(pdf_bullet(page["bullets"]))
        story.append(Spacer(1, 3)); story.append(pdf_callout(page["callout"], AMBER if page["kicker"].startswith("05") else PALE))
    story += [Spacer(1, 10), Paragraph("References and assumptions", STYLES["RGH2"])]
    refs = ["NIST SP 1339, OT Backup Quick Start Guide, June 2026. https://csrc.nist.gov/pubs/sp/1339/final", "NIST SP 800-34 Rev. 1, Contingency Planning Guide for Federal Information Systems. https://csrc.nist.gov/pubs/sp/800/34/r1/upd1/final", "NIST Cybersecurity Framework 2.0. https://www.nist.gov/cyberframework", "NATO 2026 Resilience Baseline Requirements, 15 September 2026. https://www.nato.int/en/about-us/official-texts-and-resources/official-texts/2026/09/15/2026-resilience-baseline-requirements", "Rubrik Cyber Recovery. https://www.rubrik.com/products/cyber-recovery", "Veeam Recovery Orchestrator overview. https://helpcenter.veeam.com/docs/vro/userguide/overview.html", "Cohesity Clean Room. https://www.cohesity.com/solutions/clean-room/", "MinIO upstream repository status. https://github.com/minio/minio"]
    story.append(pdf_bullet(refs)); story += [Spacer(1, 5), Paragraph("Interpretation", STYLES["RGH2"]), Paragraph("Standards and vendor pages establish context, not proof of product compliance or competitive superiority. Adoption, customer value, price, implementation effort and financial scenarios are hypotheses for validation. Real deployment requires threat modeling, privacy/security review, customer-owned recovery policy and integrations tested in a representative isolated environment.", STYLES["RGBody"]), Paragraph("Illustrative first-year unit assumptions", STYLES["RGH2"]), Paragraph("Implementation $18,000 + annual license $48,000 + annual support $12,000; direct delivery/support cost $18,000 per customer. These are planning assumptions, not market-verified prices or quotes.", STYLES["RGBody"])]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def business_pages():
    return [
        ("Opportunity and buyer", [
            "Critical operators buy backup products, infrastructure and incident response but still need a cross-system view of whether recovery can succeed. ResilienceGuard would sit above those tools as an assurance layer: inventory, dependency context, scenario assessment and evidence-backed exercises. It is intended to complement a current protection stack.",
            "A likely buyer is a leader responsible for technology (CIO), information security (CISO), operations (COO), risk or continuity. Day-to-day users may work in disaster recovery, backup, security, applications or infrastructure. An initial customer could be a regional hospital group, municipal utility, public service or critical supplier with scattered recovery records and an upcoming exercise.",
            "Deployment choices: customer-hosted on-premises/private cloud for sensitive environments; isolated managed service where data policy allows; or partner-delivered readiness assessment and exercise. Implementation work may cover service catalog mapping, recovery objective normalization, dependency workshops, adapter configuration and a supervised restore rehearsal. Ongoing revenue could include annual licensing, support, reporting and recurring exercises.",
        ], ["Customer value is measured from the customer's own data: stale/missing recovery points, blocked dependencies, observed restore duration, repeat exercise coverage and evidence preparation time.", "ROI equation: avoided outage hours × customer-validated loss/hour + validated recovery/audit labor saved − license − implementation. Do not substitute generic market averages for customer values.", "Differentiation is a hypothesis: vendor-independent scenario readiness across existing products, including simultaneous cyber and physical disruption. Vendors already sell overlapping capabilities."]),
        ("Illustrative first-year unit economics", [
            "The following are planning assumptions for discussion, not researched market prices, vendor quotes, customer commitments or a forecast. Each first-year customer pays an assumed $18,000 implementation fee, $48,000 annual license and $12,000 annual support: $78,000 first-year revenue. Direct implementation/delivery is assumed at $7,200 and recurring support/hosting at $10,800, or $18,000 total direct cost; resulting gross profit is $60,000 per customer.",
            "Overall net profitability subtracts fixed operating overhead for product engineering, sales, administration, security and legal. The illustrative overhead grows by scale in the table below. Taxes, financing, working capital, churn, customer acquisition cost and renewal rates are excluded; therefore this is a sensitivity sketch, not a company valuation or operating plan.",
        ], ["Recurring revenue mix per customer: $48,000 license + $12,000 support each year, subject to renewal.", "Implementation is one-time and should be scoped by integration count and exercise complexity.", "Direct costs rise with vendor adapters, data residency, customer-specific security controls and managed service commitments."]),
        ("Revenue scenarios and validation", [
            "The customer counts are annual first-year cohort scenarios. Revenue is cohort count multiplied by $78,000. Direct cost is cohort count multiplied by $18,000. Gross profit is revenue less direct cost. Net profitability is gross profit less assumed overhead. Numbers are USD and rounded to whole dollars.",
            "At small scale this model is loss-making because a product/security team is a fixed cost. At ten customers it reaches a small positive illustrative result; at twenty-five it assumes delivery can scale without proportional fixed cost expansion. In practice, sales cycles, procurement, assurance obligations, integration effort, support load and deployment requirements could materially change every input.",
            "Before investment or price commitments, validate buyer/problem interviews, pilot conversion, time-to-integrate, delivery hours, support effort, legal/security review requirements, renewal willingness and customer-specific outage/labor baselines. A paid pilot can be priced as a scoped implementation with an agreed evidence outcome rather than sold as guaranteed risk reduction.",
        ], ["Revenue includes implementation and annual license/support in year one; it is not ARR.", "Gross margin shown excludes fixed product, sales, compliance and corporate overhead.", "Net result is pre-tax and excludes financing, depreciation and CAC amortization."]),
    ]


def make_business_pdf():
    doc = SimpleDocTemplate(str(ROOT / "ResilienceGuard_Business_Case.pdf"), pagesize=letter, rightMargin=.68 * inch, leftMargin=.68 * inch, topMargin=.62 * inch, bottomMargin=.62 * inch, title="ResilienceGuard Business Case")
    story = [Paragraph("ResilienceGuard", STYLES["RGTitle"]), Paragraph("Business Case  /  Illustrative only", STYLES["RGH2"]), Paragraph("Commercial opportunity, buyer hypotheses and first-year revenue scenarios. Assumptions are not market-verified prices.", STYLES["RGBody"]), Paragraph("<b>How to read the numbers:</b> revenue is money received; direct costs are the delivery costs tied to customers; gross profit is revenue minus those direct costs; net result also subtracts the assumed fixed business overhead. Figures are before tax.", STYLES["RGBody"])]
    for idx, (heading, paragraphs, bullets) in enumerate(business_pages()):
        if idx: story.append(PageBreak())
        story.append(Paragraph(f"0{idx + 1}  /  {heading}", STYLES["RGH1"]))
        story.extend(Paragraph(text, STYLES["RGBody"]) for text in paragraphs)
        story.append(pdf_bullet(bullets))
        if idx == 1:
            story.append(Spacer(1, 7))
            table_rows = [["Customers", "Revenue", "Direct costs", "Gross profit", "Fixed overhead", "Net result"],
                ["1", "$78,000", "$18,000", "$60,000", "$260,000", "-$200,000"],
                ["5", "$390,000", "$90,000", "$300,000", "$360,000", "-$60,000"],
                ["10", "$780,000", "$180,000", "$600,000", "$520,000", "$80,000"],
                ["25", "$1,950,000", "$450,000", "$1,500,000", "$750,000", "$750,000"]]
            table = Table(table_rows, colWidths=[.69*inch, 1.14*inch, 1.06*inch, 1.1*inch, 1.19*inch, 1.1*inch], repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7), ("LEADING", (0, 0), (-1, -1), 9), ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#D8E2E6")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6F7")]), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
            story.append(table); story.append(Spacer(1, 6)); story.append(Paragraph("Net result = gross profit − fixed overhead. Year-one figures include the assumed implementation fee and one annual license/support contract. All figures are illustrative USD.", STYLES["RGSmall"]))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def make_evaluation_pdf():
    doc = SimpleDocTemplate(str(ROOT / "ResilienceGuard_Evaluation_Report.pdf"), pagesize=letter, rightMargin=.68 * inch, leftMargin=.68 * inch, topMargin=.62 * inch, bottomMargin=.62 * inch, title="ResilienceGuard Evaluation Report")
    story = [Paragraph("ResilienceGuard", STYLES["RGTitle"]), Paragraph("Evaluation Report  /  Synthetic experiment", STYLES["RGH2"]), Paragraph(f"Fixed seed {evaluation.get('seed', 'not recorded')} · reference {evaluation.get('reference_time', 'not recorded')} · no customer or production data", STYLES["RGBody"]), Paragraph("What was evaluated", STYLES["RGH1"]), Paragraph("Automated tests and a repeatable synthetic evaluation exercise whether backup files match their saved fingerprints, how old they are, whether service requirements are missing, and whether an approved sample file can be copied into an isolated folder. Three actual local copies were timed; those measurements cover file copying and sample-file checks only, not restoration of an application or database.", STYLES["RGBody"]), Paragraph("How to read the results: each fraction is the number of synthetic examples classified correctly divided by the number checked. For example, 5/5 means five of five examples matched their expected label. It does not mean five real services are safe or that real-world recovery is 100% reliable.", STYLES["RGBody"])]
    suite = [["Automated suite", "Passed", "Failed", "Total"], ["pytest backend/tests", str(pytest_summary.get("passed", "not recorded")), str(pytest_summary.get("failed", "not recorded")), str(pytest_summary.get("total", "not recorded"))]]
    t = Table(suite, colWidths=[3.1*inch, 1.0*inch, 1.0*inch, 1.0*inch]); t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#"+NAVY)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#D5E0E4")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#EFF5F6")]), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)])); story += [t, Spacer(1, 13), Paragraph("Measured checks", STYLES["RGH1"])]
    check_rows = [["Check", "Correct / total", "Observed"]]
    for row in evaluation.get("checks", []):
            plain_names = {
                "SHA-256 expected good/bad classification": "File fingerprint: expected match / mismatch",
                "RPO freshness classification": "Backup age compared with data-loss target",
                "Missing backup record classification": "Services with missing backup records",
                "Recovery dependency matrix (4 services × 3 scenarios)": "Service blockers across 3 scenarios",
                "Approved isolated restore copy + hash + fixture health": "Approved test copy, fingerprint, and file check",
                "Unauthorized/untrusted restore prevention": "Pending or untrusted copy was denied",
                "Evidence/action grounding against rule-only baseline": "Assistant references matched allowed evidence and rules",
            }
            check_rows.append([plain_names.get(row["name"], row["name"]), f"{row['passed']} / {row['total']}", f"{row['rate']*100:.1f}%"])
    t = Table(check_rows, colWidths=[4.25*inch, 1.3*inch, 1.0*inch], repeatRows=1); t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#"+TEAL)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.6), ("LEADING", (0, 0), (-1, -1), 10), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#D8E2E6")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7F8")]), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)])); story.append(t)
    story += [PageBreak(), Paragraph("Restore timing and AI baseline", STYLES["RGH1"]), Paragraph(f"Approved file restore samples: {restore_stats.get('samples', 0)}. Mean observed time: {fmt_seconds(restore_stats.get('mean'))}; minimum: {fmt_seconds(restore_stats.get('minimum'))}; maximum: {fmt_seconds(restore_stats.get('maximum'))}.", STYLES["RGBody"]), Paragraph(restore_stats.get("measurement_scope", "No restore timing was recorded."), STYLES["RGSmall"]), Paragraph(f"Rule-only baseline action IDs: {', '.join(evaluation.get('rule_only_baseline_action_ids', [])) or 'none'}. Assistant mode: {evaluation.get('assistant_mode', 'not recorded')}. Assistant action IDs: {', '.join(evaluation.get('assistant_action_ids', [])) or 'none'}. Plan action parity: {evaluation.get('plan_action_parity', 'not recorded')}.", STYLES["RGBody"]), Paragraph(f"Evidence IDs and allowed action IDs were checked against the synthetic assessment: {fmt_pct(checks.get('Evidence/action grounding against rule-only baseline', {'passed': 0, 'total': 0, 'rate': 0}))}. The evaluated assistant used the deterministic fallback. {evaluation.get('local_llm_accuracy', 'Local model performance not recorded')}.", STYLES["RGBody"]), Paragraph("Security dependency scans", STYLES["RGH1"])]
    pysec = security.get("python", {}); jssec = security.get("npm", {})
    story.append(Paragraph(f"pip-audit: {pysec.get('summary', 'result not recorded')}. npm audit: {jssec.get('summary', 'result not recorded')}. The PptxGenJS tree pins image-size 2.0.3 through an npm override after the initial scan found high-severity image-size parser advisories. The lockfile was re-audited after that fix. Results are point-in-time dependency scans, not a complete code/container penetration test.", STYLES["RGBody"]))
    story += [Spacer(1, 5), pdf_callout("The evaluation does not claim malware cleanliness, production recoverability, clinical application health, live PostgreSQL restore health, or generative-model accuracy.", AMBER), Spacer(1, 13), Paragraph("Interpretation and limits", STYLES["RGH1"]), Paragraph("All fixtures are deterministic synthetic JSON with a fixed timestamp. Known expected labels are encoded in the evaluation script and matched against engine outputs. The sample size is deliberately small and tests correctness of the MVP's explicit rules, not performance at scale. The restoration test uses local filesystem copy operations; its measured milliseconds cannot be compared with a database or service RTO. The security evaluation is a dependency advisory check; it is not a full static/dynamic security assessment.", STYLES["RGBody"]), Paragraph("Unmeasured follow-up", STYLES["RGH2"]), pdf_bullet(["Run Qwen3 on a named hardware/runtime profile and score claims, evidence references, uncertainties and divergence from the rule-only baseline.", "Validate recovery plans with backup/vendor adapters and customer-owned failure-domain inventories.", "Exercise malware scanning, isolated application/database health checks, authentic identity/approval controls and independent audit protection."])]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    ART.mkdir(parents=True, exist_ok=True)
    make_idea_docx()
    make_idea_pdf()
    make_business_pdf()
    make_evaluation_pdf()
    print("Generated Idea DOCX/PDF, Business Case PDF, and Evaluation Report PDF.")
